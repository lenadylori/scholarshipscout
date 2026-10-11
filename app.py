import os
from io import BytesIO
from pathlib import Path

import streamlit as st

st.set_page_config(page_title="Scholarship Scout | Find Your Future", page_icon="✦", layout="wide", initial_sidebar_state="expanded")

def secret(name):
    try:
        return st.secrets.get(name, os.getenv(name, ""))
    except Exception:
        return os.getenv(name, "")

# Editorial, layered teal-and-ivory visual system inspired by the supplied references.
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Manrope:wght@400;500;600;700;800&display=swap');
:root { --deep:#064e49; --deep2:#003d39; --teal:#08766c; --mint:#d8e8df; --paper:#eef1e9; --ink:#063f3b; --muted:#54736d; --line:rgba(5,70,64,.28); }
html, body, [class*="css"] { font-family:'DM Sans',sans-serif; }
.stApp { background: var(--paper); color:var(--ink); }
[data-testid="stHeader"] { background:transparent; }
.block-container { max-width: 1240px; padding-top: 1.5rem; padding-bottom: 4rem; }
h1,h2,h3,h4 { font-family:'Manrope',sans-serif !important; color:var(--ink) !important; letter-spacing:-.035em; }
h1 { font-size:clamp(2.6rem,5vw,5rem) !important; line-height:1.0 !important; font-weight:800 !important; }
h2 { font-weight:800 !important; }
p, label, .stMarkdown { color:var(--ink); }
div[data-testid="stSidebar"] { background:#dce7df; border-right:1px solid var(--line); }
div[data-testid="stSidebar"] h2, div[data-testid="stSidebar"] label { color:var(--ink); }
div[data-testid="stSidebar"] [data-testid="stVerticalBlock"] { gap:.7rem; }
div[data-testid="stTextInput"] input, div[data-testid="stTextArea"] textarea, div[data-testid="stSelectbox"] div[data-baseweb="select"], div[data-testid="stNumberInput"] input {
  border-radius:14px !important; background:#f8f8f1 !important; border-color:var(--line) !important; color:var(--ink) !important;
}
.stButton > button, .stDownloadButton > button, .stLinkButton > a {
  border-radius:999px !important; border:1px solid var(--deep) !important; background:var(--deep) !important; color:#f1f3eb !important;
  font-weight:700 !important; padding:.65rem 1.2rem !important; transition:transform .15s ease, background .15s ease;
}
.stButton > button:hover, .stDownloadButton > button:hover, .stLinkButton > a:hover { background:#0a6b61 !important; transform:translateY(-1px); }
button[kind="secondary"] { background:transparent !important; color:var(--deep) !important; }
[data-testid="stTabs"] [data-baseweb="tab-list"] { gap:.5rem; background:#dbe6dd; padding:.45rem; border-radius:999px; }
[data-testid="stTabs"] button[role="tab"] { border-radius:999px; color:var(--ink); padding:.65rem 1.1rem; }
[data-testid="stTabs"] button[aria-selected="true"] { background:var(--deep) !important; color:#f4f4eb !important; }
[data-testid="stVerticalBlockBorderWrapper"] { background:rgba(250,250,244,.74); border:1px solid var(--line); border-radius:24px; box-shadow:0 10px 28px rgba(0,54,49,.06); }
[data-testid="stAlert"] { border-radius:18px; }
div[data-testid="stMetric"] { background:#dbe7df; padding:1rem; border-radius:18px; border:1px solid var(--line); }
.hero { position:relative; overflow:hidden; padding:clamp(1.5rem,4vw,3.5rem); border-radius:34px; color:#f2f2e8; background:radial-gradient(circle at 85% 12%,rgba(121,190,166,.25),transparent 28%),linear-gradient(135deg,#003d39 0%,#075e56 60%,#0b7167 100%); box-shadow:0 20px 45px rgba(0,53,49,.18); margin-bottom:1.3rem; }
.hero:after { content:"✦"; position:absolute; right:5%; bottom:-65px; font-size:240px; line-height:1; color:rgba(224,239,226,.08); transform:rotate(-12deg); pointer-events:none; }
.eyebrow { font-size:.76rem; font-weight:800; letter-spacing:.22em; text-transform:uppercase; color:#a9d6c5; margin-bottom:1rem; }
.hero h1 { color:#f2f2e8 !important; max-width:850px; }
.hero p { color:#d0e3d8; max-width:650px; font-size:1.05rem; line-height:1.65; }
.hero-pill { display:inline-block; border:1px solid rgba(229,241,230,.5); border-radius:999px; padding:.4rem .85rem; margin:.25rem .25rem 0 0; color:#edf2e8; font-size:.8rem; }
.section-kicker { color:#0a776c; font-size:.75rem; font-weight:800; letter-spacing:.18em; text-transform:uppercase; }
.panel-title { font-size:1.35rem; font-weight:800; letter-spacing:-.04em; }
.result-url { font-size:.82rem; color:#0b766b; overflow-wrap:anywhere; }
.small-note { color:#54736d; font-size:.85rem; }
hr { border-color:var(--line) !important; }
footer {visibility:hidden;}
@media (max-width: 700px) {
 .block-container { padding-left:1rem; padding-right:1rem; }
 .hero { border-radius:24px; }
 h1 { font-size:2.7rem !important; }
}
</style>
""", unsafe_allow_html=True)

with st.sidebar:
    st.markdown("## ✦ Scholarship Scout")
    st.caption("FIND YOUR FUTURE")
    st.markdown("---")
    st.markdown("### Your search")
    location = st.text_input("Location", placeholder="New Jersey, United States")
    education = st.selectbox("Education level", ["High school senior", "High school student", "Undergraduate", "Graduate", "Trade / vocational", "Other"])
    field = st.text_input("Major or career", placeholder="Biomedical engineering")
    graduation_year = st.text_input("Graduation year", placeholder="2027")
    min_award = st.number_input("Minimum award ($)", min_value=0, value=0, step=500)
    max_award = st.number_input("Maximum award ($, 0 = any)", min_value=0, value=0, step=1000)
    st.markdown("---")
    st.markdown("### Optional eligibility themes")
    st.caption("Use only details you're comfortable sharing.")
    ethnicity = st.text_input("Ethnicity / heritage", placeholder="Only if relevant")
    lgbtq = st.selectbox("LGBTQ+ scholarships", ["No preference", "Include relevant scholarships", "Prioritize relevant scholarships"])
    disability = st.selectbox("Disability / neurodiversity", ["No preference", "Include relevant scholarships", "Prioritize relevant scholarships"])
    adversity = st.multiselect("Experiences and interests", [
        "Adversity / overcoming obstacles", "Family or financial hardship", "Trauma or recovery",
        "Disability or chronic barriers", "LGBTQ+ experiences", "First-generation student",
        "Community service", "Leadership", "Academic persistence", "Women in STEM"
    ])
    struggles = st.text_area("Other keywords", placeholder="Optional search terms, no need to share details.")
    count = st.slider("Results to find", 5, 20, 10)
    st.markdown("---")
    st.caption("✦ Your story belongs to you. Review every answer before using it.")

st.markdown("""
<div class="hero">
  <div class="eyebrow">Scholarship discovery · Personal essay studio</div>
  <h1>Your story.<br> Your next opportunity.</h1>
  <p>Find scholarships that fit your goals and experiences. Turn your own notes and writing into thoughtful, personal application drafts without losing your voice.</p>
  <span class="hero-pill">✦ Personalized search</span>
  <span class="hero-pill">✦ Your voice, not a template</span>
  <span class="hero-pill">✦ One organized workspace</span>
</div>
""", unsafe_allow_html=True)

st.markdown('<div class="section-kicker">YOUR WORKSPACE</div>', unsafe_allow_html=True)
tab_search, tab_essay, tab_tracker, tab_resources = st.tabs(["✦ Discover", "✎ Essay studio", "▤ My checklist", "↗ Resources"])

def search_web(query, limit=10):
    key = secret("TAVILY_API_KEY")
    if not key:
        return None, "Add TAVILY_API_KEY in Streamlit Secrets to enable live internet search."
    try:
        from tavily import TavilyClient
        client = TavilyClient(api_key=key)
        response = client.search(query=query, search_depth="advanced", max_results=limit, include_answer=False)
        return response.get("results", []), None
    except Exception as e:
        return None, f"Search failed: {e}"

def make_search_query():
    parts = [
        "scholarships official application",
        location, education, field, graduation_year, ethnicity,
        "LGBTQ scholarship" if lgbtq != "No preference" else "",
        "disability neurodiversity scholarship" if disability != "No preference" else "",
        " ".join(adversity), struggles,
        f"scholarship award ${min_award} or more" if min_award else "",
        f"scholarship maximum ${max_award}" if max_award else "",
    ]
    return " ".join(p.strip() for p in parts if p and p.strip())

def extract_document(uploaded):
    if uploaded is None:
        return ""
    suffix = Path(uploaded.name).suffix.lower()
    data = uploaded.getvalue()
    try:
        if suffix in [".txt", ".md"]:
            return data.decode("utf-8", errors="replace")[:30000]
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data))
            return "\n".join(page.extract_text() or "" for page in reader.pages)[:30000]
        if suffix == ".docx":
            from docx import Document
            doc = Document(BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs)[:30000]
    except Exception as e:
        st.warning(f"Could not read {uploaded.name}: {e}")
    return ""

def draft_with_gpt(prompt, source_text, profile_text, scholarship_text):
    key = secret("OPENAI_API_KEY")
    if not key:
        return None, "Add OPENAI_API_KEY in Streamlit Secrets to enable AI essay drafting."
    try:
        from openai import OpenAI
        client = OpenAI(api_key=key)
        instructions = (
            "You are a thoughtful scholarship essay coach. Help the applicant present real experiences clearly, "
            "specifically, and compellingly. Do not invent facts, awards, diagnoses, identities, hardships, or outcomes. "
            "Do not exaggerate trauma. If details are missing, use [ADD YOUR DETAIL] placeholders. Match the exact prompt "
            "and word limit. Preserve the applicant's voice based on their supplied writing sample. Explain how the draft "
            "answers the selection criteria, but never promise or guarantee an award."
        )
        user_content = f"""SCHOLARSHIP PROMPT / REQUIREMENTS:
{scholarship_text}

APPLICANT PROFILE (user-provided):
{profile_text}

APPLICANT'S WRITING SAMPLE / NOTES:
{source_text}

TASK:
{prompt}

Write a strong, authentic draft grounded only in supplied facts. After the draft, list facts to verify and briefly explain how the response aligns with the stated criteria. Do not include irrelevant private information."""
        response = client.responses.create(model="gpt-4.1-mini", instructions=instructions, input=user_content)
        return response.output_text, None
    except Exception as e:
        return None, f"Drafting failed: {e}"

with tab_search:
    left, right = st.columns([1.45, 1], gap="large")
    with left:
        st.markdown('<div class="section-kicker">01 / MATCH WITH PURPOSE</div>', unsafe_allow_html=True)
        st.markdown("## Find opportunities that fit")
        st.write("Search scholarship pages using your goals and chosen eligibility themes. Verify deadlines and criteria with the official sponsor.")
        if st.button("Search scholarships ↗", type="primary", key="search_button"):
            query = make_search_query()
            st.session_state["last_query"] = query
            with st.spinner("Searching the web for scholarship leads..."):
                results, error = search_web(query, count)
            if error:
                st.warning(error)
            else:
                st.session_state["scholarship_results"] = results
                st.success(f"Found {len(results)} search results.")
        if st.session_state.get("last_query"):
            st.caption("Current search: " + st.session_state["last_query"])
    with right:
        with st.container(border=True):
            st.markdown('<div class="section-kicker">YOUR SEARCH SNAPSHOT</div>', unsafe_allow_html=True)
            st.markdown("### Built around you")
            st.write(f"**Education:** {education}")
            st.write(f"**Field:** {field or 'Add a major or career in the sidebar'}")
            st.write(f"**Location:** {location or 'Any location'}")
            st.write(f"**Themes:** {', '.join(adversity) if adversity else 'No experience themes selected'}")
            st.caption("Sensitive filters are optional and are used to make search queries more relevant.")
    results = st.session_state.get("scholarship_results", [])
    for i, item in enumerate(results, 1):
        with st.container(border=True):
            a, b = st.columns([5, 1])
            with a:
                st.markdown(f"### {i:02d} · {item.get('title', 'Scholarship result')}")
                st.write(item.get("content", "No preview available."))
                st.markdown(f'<div class="result-url">{item.get("url", "")}</div>', unsafe_allow_html=True)
            with b:
                if item.get("url"):
                    st.link_button("Open ↗", item["url"])
    if results:
        st.info("These are search leads, not verified awards. Confirm eligibility, award amount, deadlines, and AI-use rules on the sponsor's official website.")

with tab_essay:
    st.markdown('<div class="section-kicker">02 / WRITE IN YOUR OWN VOICE</div>', unsafe_allow_html=True)
    st.markdown("## Essay studio")
    st.write("Give the AI your prompt and a sample of your writing. It can learn your tone from the sample and use only facts you provide.")
    with st.container(border=True):
        uploaded = st.file_uploader("Upload your writing or notes", type=["txt", "md", "pdf", "docx"], key="writing_sample")
        writing_sample = extract_document(uploaded)
        if uploaded:
            st.success(f"Loaded text from {uploaded.name} ({len(writing_sample)} characters).")
            with st.expander("Preview extracted text"):
                st.text(writing_sample[:4000] or "No readable text found.")
        essay_profile = st.text_area("Facts you want the essay to draw from", placeholder="Your goals, projects, responsibilities, activities, and experiences relevant to this application.", height=130)
        scholarship_prompt = st.text_area("Scholarship prompt and criteria", placeholder="Paste the exact question, word limit, and selection criteria.", height=130)
        task = st.text_area("What should the essay emphasize?", placeholder="Connect my engineering interests and community service to this scholarship's mission.", height=80)
        consent = st.checkbox("I understand that my uploaded writing and details will be sent to the configured AI provider to generate a draft.")
        if st.button("Create my draft ✦", type="primary"):
            if not consent:
                st.error("Please confirm how your information will be processed.")
            elif not scholarship_prompt.strip():
                st.error("Add the scholarship prompt first.")
            else:
                profile_text = "\n".join([
                    f"Education: {education}", f"Location: {location}", f"Field: {field}",
                    f"Graduation year: {graduation_year}", f"Selected themes: {', '.join(adversity)}",
                    f"Other user-provided keywords: {struggles}", essay_profile
                ])
                with st.spinner("Building a draft from your details..."):
                    result, error = draft_with_gpt(task, writing_sample, profile_text, scholarship_prompt)
                if error:
                    st.warning(error)
                else:
                    st.session_state["essay_draft"] = result
    if st.session_state.get("essay_draft"):
        st.markdown("### Your draft")
        st.text_area("Edit before using", value=st.session_state["essay_draft"], height=380, key="editable_draft")
        st.download_button("Download essay draft ↓", st.session_state["editable_draft"], file_name="scholarship-essay-draft.txt", mime="text/plain")
        st.warning("Read it carefully, verify every claim, and follow the scholarship's rules about AI assistance.")

with tab_tracker:
    st.markdown('<div class="section-kicker">03 / KEEP MOVING</div>', unsafe_allow_html=True)
    st.markdown("## Application checklist")
    st.write("A simple checklist for keeping each application organized.")
    items = [
        "Confirm eligibility on the official website", "Verify deadline and time zone",
        "Read the scholarship's AI-use policy", "Request recommendation letters",
        "Prepare transcript and required documents", "Proofread each answer for accuracy",
        "Submit the application yourself", "Save the confirmation email or receipt"
    ]
    cols = st.columns(2)
    for idx, item in enumerate(items):
        with cols[idx % 2]:
            st.checkbox(item, key=f"check_{idx}")
    st.caption("This starter checklist is kept in the current session and is not yet saved to a permanent account.")

with tab_resources:
    st.markdown('<div class="section-kicker">04 / TRUSTED STARTING POINTS</div>', unsafe_allow_html=True)
    st.markdown("## Scholarship resources")
    resources = [
        ("CareerOneStop Scholarship Finder", "https://www.careeronestop.org/Toolkit/Training/find-scholarships.aspx"),
        ("Federal Student Aid", "https://studentaid.gov/"),
        ("BigFuture Scholarship Search", "https://bigfuture.collegeboard.org/pay-for-college/scholarship-search"),
        ("Scholarship America", "https://scholarshipamerica.org/students/"),
        ("Fastweb", "https://www.fastweb.com/"),
    ]
    for name, url in resources:
        with st.container(border=True):
            st.markdown(f"### {name}")
            st.link_button("Visit resource ↗", url)
    st.caption("Be cautious of scams. Legitimate scholarships should not require you to pay to apply.")

st.markdown("---")
st.markdown('<div class="small-note">SCHOLARSHIP SCOUT AI · Built to help you find opportunities and tell your story with care. No essay can guarantee an award. Review all generated content before submitting.</div>', unsafe_allow_html=True)
